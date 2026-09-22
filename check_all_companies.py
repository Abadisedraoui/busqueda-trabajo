#!/usr/bin/env python3
"""
check_all_companies.py

Monitoriza TODAS las filas de watchlist_consolidado.xlsx.

Principios:
- Nunca saltar una empresa en silencio: toda fila termina en coverage_report.csv.
- Preferir APIs/feeds públicos del ATS cuando existen.
- Si el ATS no está resuelto, inspeccionar la web oficial y detectar el ATS.
- Si la URL falta o es un agregador (LinkedIn, Indeed...), descubrir la web de empleo
  oficial mediante una búsqueda ligera y cachear el resultado.
- Para portales JS, usar Playwright como fallback opcional (--browser).
- Una oferta solo entra si el enlace final pertenece a una web oficial o ATS de empleo,
  nunca a LinkedIn/InfoJobs/Indeed/Glassdoor.

Salidas:
- nuevas_ofertas.csv              solo ofertas nuevas de diseño
- all_open_design_jobs.csv        snapshot actual de ofertas de diseño encontradas
- ofertas_vistas.csv              histórico para no repetir
- coverage_report.csv             una fila por empresa del Excel, sin skips silenciosos
- resolution_cache.json           URLs/ATS descubiertos para reutilizar al día siguiente
- daily_search_report.md          resumen legible de la ejecución

Uso:
    python check_all_companies.py --browser
    python check_all_companies.py --company "Notion" --browser
    python check_all_companies.py --max-companies 20 --browser

Dependencias:
    pip install -r requirements.txt
    playwright install chromium   # solo si se usa --browser
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html as html_lib
import json
import re
import sys
import time
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional
from urllib.parse import parse_qs, quote_plus, urlencode, urljoin, urlparse, urlunparse

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from bs4 import BeautifulSoup
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
WATCHLIST = BASE / "watchlist_consolidado.xlsx"
SEEN_CSV = BASE / "ofertas_vistas.csv"
NEW_CSV = BASE / "nuevas_ofertas.csv"
OPEN_CSV = BASE / "all_open_design_jobs.csv"
COVERAGE_CSV = BASE / "coverage_report.csv"
CACHE_JSON = BASE / "resolution_cache.json"
REPORT_MD = BASE / "daily_search_report.md"
DASHBOARD_HTML = BASE / "dashboard.html"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36 "
    "JobWatch/2.0"
)

TIMEOUT = 20
REQUEST_DELAY = 0.12
MAX_GENERIC_PAGES = 6
CACHE_MAX_AGE_DAYS = 30

# Se busca de forma deliberadamente amplia: la fase de puntuación decide después si
# Web/UI/Visual/etc. encaja realmente. Aquí el objetivo es no perder oportunidades.
DESIGN_TITLE_PATTERNS = [
    r"\bproduct\s+(?:ux[ /-]?ui\s+)?designer\b",
    r"\bsenior\s+product\s+designer\b",
    r"\bdigital\s+product\s+designer\b",
    r"\bux\s*(?:/|&|and|-)?\s*ui\s+designer\b",
    r"\bui\s*(?:/|&|and|-)?\s*ux\s+designer\b",
    r"\bux\s+designer\b",
    r"\bui\s+designer\b",
    r"\bweb\s+designer\b",
    r"\bweb\s+ux\s+designer\b",
    r"\bdigital\s+designer\b",
    r"\bvisual\s+designer\b",
    r"\bproduct\s+visual\s+designer\b",
    r"\buser\s+experience\s+designer\b",
    r"\buser\s+interface\s+designer\b",
    r"\bexperience\s+designer\b",
    r"\binteraction\s+designer\b",
    r"\binterface\s+designer\b",
    r"\bdesign\s+systems?\s+designer\b",
    r"\bdesign\s+systems?\b",
    r"\bproduct\s+design\b",
    r"\bux\s+design\b",
    r"\bui\s+design\b",
    r"\bdisenador(?:a|/a)?\s+(?:de\s+)?producto\b",
    r"\bdiseno\s+de\s+producto\b",
    r"\bdisenador(?:a|/a)?\s+ux\b",
    r"\bdisenador(?:a|/a)?\s+ui\b",
    r"\bdisenador(?:a|/a)?\s+web\b",
]

# Evita falsos positivos obvios de diseño físico/ingeniería provocados por
# expresiones como "Product Design Engineer". No excluye management/head: esos
# roles se encuentran y luego se puntúan/descartan según las reglas de encaje.
NON_DIGITAL_TITLE_PATTERNS = [
    r"\bproduct\s+design\s+engineer\b",
    r"\bmechanical\b", r"\bindustrial\s+designer\b",
    r"\binterior\s+designer\b", r"\bfashion\s+designer\b",
    r"\bcad\b", r"\bhardware\s+designer\b",
]

TRACKING_QUERY_PREFIXES = ("utm_",)
TRACKING_QUERY_KEYS = {
    "source", "ref", "referrer", "trackingid", "trk", "gh_src", "lever-source",
    "campaign", "campaignid", "fbclid", "gclid", "mc_cid", "mc_eid",
}

AGGREGATOR_DOMAINS = {
    "linkedin.com", "www.linkedin.com", "indeed.com", "www.indeed.com",
    "glassdoor.com", "www.glassdoor.com", "infojobs.net", "www.infojobs.net",
    "jobgether.com", "www.jobgether.com", "jooble.org", "www.jooble.org",
    "talent.com", "www.talent.com", "simplyhired.com", "www.simplyhired.com",
}

ATS_DOMAINS = {
    "greenhouse": ("greenhouse.io",),
    "lever": ("lever.co",),
    "ashby": ("ashbyhq.com",),
    "workable": ("workable.com",),
    "workday": ("myworkdayjobs.com", "myworkdaysite.com"),
    "smartrecruiters": ("smartrecruiters.com",),
    "recruitee": ("recruitee.com",),
    "personio": ("jobs.personio.de", "jobs.personio.com"),
    "icims": ("icims.com",),
    "jobvite": ("jobvite.com",),
    "breezyhr": ("breezy.hr",),
}

KNOWN_PLATFORM_ALIASES = {
    "no_se": "unknown",
    "verificar": "unknown",
    "otro": "unknown",
    "agencia_excluir": "unknown",
    "": "unknown",
}


@dataclass
class CompanyRow:
    row_number: int
    company: str
    input_url: str
    platform_hint: str
    notes: str


@dataclass
class Job:
    company: str
    title: str
    url: str
    platform: str
    location: str = ""
    description: str = ""
    external_id: str = ""

    @property
    def canonical_url(self) -> str:
        return canonicalize_url(self.url)

    @property
    def signature(self) -> str:
        raw = f"{norm(self.company)}|{norm(self.title)}|{self.external_id or self.canonical_url}"
        return hashlib.sha1(raw.encode("utf-8")).hexdigest()

    @property
    def content_fingerprint(self) -> str:
        """Detecta requisiciones duplicadas con IDs/URLs diferentes pero mismo contenido."""
        desc = norm(self.description)
        if len(desc) < 120:
            return ""
        raw = f"{norm(self.company)}|{norm(self.title)}|{norm(self.location)}|{desc[:6000]}"
        return hashlib.sha1(raw.encode("utf-8")).hexdigest()


class BrowserFetcher:
    def __init__(self, enabled: bool):
        self.enabled = enabled
        self._pw = None
        self._browser = None
        self._context = None

    def _ensure(self):
        if not self.enabled:
            return False
        if self._browser:
            return True
        try:
            from playwright.sync_api import sync_playwright
            self._pw = sync_playwright().start()
            self._browser = self._pw.chromium.launch(headless=True)
            self._context = self._browser.new_context(user_agent=USER_AGENT)
            # No necesitamos imágenes, vídeo ni fuentes para detectar ATS/ofertas.
            # Bloquearlos acelera mucho una ejecución de cientos de empresas.
            def _route(route):
                if route.request.resource_type in {"image", "media", "font"}:
                    route.abort()
                else:
                    route.continue_()
            self._context.route("**/*", _route)
            return True
        except Exception as exc:
            print(f"[browser] No disponible: {exc}", file=sys.stderr)
            self.enabled = False
            return False

    def fetch(self, url: str) -> tuple[str, str]:
        if not self._ensure():
            return "", url
        page = self._context.new_page()
        try:
            last_exc = None
            for attempt in range(2):
                try:
                    page.goto(url, wait_until="domcontentloaded", timeout=25000)
                    page.wait_for_timeout(1400 + attempt * 700)
                    return page.content(), page.url
                except Exception as exc:
                    last_exc = exc
                    if attempt == 0:
                        page.wait_for_timeout(800)
            raise last_exc
        finally:
            page.close()

    def close(self):
        try:
            if self._context:
                self._context.close()
            if self._browser:
                self._browser.close()
            if self._pw:
                self._pw.stop()
        except Exception:
            pass


def norm(value: str) -> str:
    value = str(value or "").strip().lower()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(c for c in value if not unicodedata.combining(c))
    value = re.sub(r"\s+", " ", value)
    return value


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", BeautifulSoup(value or "", "html.parser").get_text(" ")).strip()


def canonicalize_url(url: str) -> str:
    """Normaliza sin borrar parámetros que pueden identificar una oferta.

    Antes se eliminaba toda la query string; eso podía convertir dos ofertas distintas
    en la misma URL en portales que usan ?jobId=... . Ahora solo quitamos tracking.
    """
    if not url:
        return ""
    try:
        p = urlparse(url.strip())
        host = p.netloc.lower().replace(":443", "")
        path = re.sub(r"/+", "/", p.path or "/")
        path = path.rstrip("/") or "/"
        kept = []
        for key, values in parse_qs(p.query, keep_blank_values=True).items():
            kl = key.lower()
            if kl in TRACKING_QUERY_KEYS or any(kl.startswith(x) for x in TRACKING_QUERY_PREFIXES):
                continue
            for value in values:
                kept.append((key, value))
        query = urlencode(sorted(kept), doseq=True)
        return urlunparse((p.scheme.lower() or "https", host, path, "", query, ""))
    except Exception:
        return url.strip()


def host_of(url: str) -> str:
    try:
        return urlparse(url).netloc.lower().split(":")[0]
    except Exception:
        return ""


def is_aggregator(url: str) -> bool:
    host = host_of(url)
    return any(host == d or host.endswith("." + d) for d in AGGREGATOR_DOMAINS)


def is_design_title(title: str) -> bool:
    t = norm(title)
    if not t:
        return False
    if any(re.search(p, t, re.I) for p in NON_DIGITAL_TITLE_PATTERNS):
        return False
    return any(re.search(p, t, re.I) for p in DESIGN_TITLE_PATTERNS)


def build_session() -> requests.Session:
    session = requests.Session()
    retry = Retry(
        total=3, connect=3, read=2,
        backoff_factor=0.7,
        status_forcelist=(408, 425, 429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET", "HEAD", "OPTIONS", "POST"}),
        respect_retry_after_header=True,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_connections=20, pool_maxsize=20)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def session_get(session: requests.Session, url: str, **kwargs) -> requests.Response:
    headers = dict(kwargs.pop("headers", {}))
    headers.setdefault("User-Agent", USER_AGENT)
    headers.setdefault("Accept-Language", "en-US,en;q=0.9,es;q=0.8")
    r = session.get(url, headers=headers, timeout=TIMEOUT, allow_redirects=True, **kwargs)
    r.raise_for_status()
    return r


def session_post(session: requests.Session, url: str, **kwargs) -> requests.Response:
    headers = dict(kwargs.pop("headers", {}))
    headers.setdefault("User-Agent", USER_AGENT)
    headers.setdefault("Accept", "application/json,text/plain,*/*")
    r = session.post(url, headers=headers, timeout=TIMEOUT, allow_redirects=True, **kwargs)
    r.raise_for_status()
    return r


def load_watchlist(path: Path) -> list[CompanyRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["watchlist"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
    required = {"empresa", "web_empleo", "plataforma_ats"}
    missing = required - set(idx)
    if missing:
        raise RuntimeError(f"Faltan columnas en watchlist: {sorted(missing)}")
    rows: list[CompanyRow] = []
    for n, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        company = values[idx["empresa"]] if idx["empresa"] < len(values) else None
        if not company:
            continue
        url = values[idx["web_empleo"]] if idx["web_empleo"] < len(values) else ""
        platform = values[idx["plataforma_ats"]] if idx["plataforma_ats"] < len(values) else ""
        notes = values[idx.get("notas", -1)] if "notas" in idx and idx["notas"] < len(values) else ""
        rows.append(CompanyRow(n, str(company).strip(), str(url or "").strip(), str(platform or "").strip().lower(), str(notes or "").strip()))
    return rows


def load_cache() -> dict:
    if not CACHE_JSON.exists():
        return {}
    try:
        return json.loads(CACHE_JSON.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_cache(cache: dict):
    CACHE_JSON.write_text(json.dumps(cache, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")


def cache_is_fresh(entry: dict) -> bool:
    stamp = (entry or {}).get("last_success", "")
    if not stamp:
        return False
    try:
        age = datetime.now() - datetime.fromisoformat(stamp)
        return age.days <= CACHE_MAX_AGE_DAYS
    except Exception:
        return False


def ensure_seen_schema():
    headers = ["empresa", "titulo", "url", "fecha_visto", "signature", "content_fingerprint"]
    if not SEEN_CSV.exists():
        return headers
    with SEEN_CSV.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        old_headers = reader.fieldnames or []
        rows = list(reader)
    if old_headers == headers:
        return headers
    tmp = SEEN_CSV.with_suffix(".tmp.csv")
    with tmp.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        for row in rows:
            w.writerow({h: row.get(h, "") for h in headers})
    tmp.replace(SEEN_CSV)
    return headers


def load_seen() -> tuple[set[str], set[str], set[str]]:
    urls: set[str] = set()
    signatures: set[str] = set()
    fingerprints: set[str] = set()
    if not SEEN_CSV.exists():
        return urls, signatures, fingerprints
    ensure_seen_schema()
    with SEEN_CSV.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row.get("url"):
                urls.add(canonicalize_url(row["url"]))
            if row.get("signature"):
                signatures.add(row["signature"])
            if row.get("content_fingerprint"):
                fingerprints.add(row["content_fingerprint"])
    return urls, signatures, fingerprints


def append_seen(jobs: list[Job]):
    headers = ensure_seen_schema()
    exists = SEEN_CSV.exists()
    with SEEN_CSV.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        if not exists:
            w.writeheader()
        for j in jobs:
            w.writerow({
                "empresa": j.company,
                "titulo": j.title,
                "url": j.url,
                "fecha_visto": datetime.now().strftime("%Y-%m-%d"),
                "signature": j.signature,
                "content_fingerprint": j.content_fingerprint,
            })

def detect_platform_from_url(url: str) -> str:
    h = host_of(url)
    for platform, domains in ATS_DOMAINS.items():
        if any(h == d or h.endswith("." + d) or d in h for d in domains):
            return platform
    if "teamtailor" in h:
        return "teamtailor"
    return ""


def detect_platform_from_html(html: str, base_url: str = "") -> str:
    """Detecta ATS incluso cuando usa un dominio personalizado/CNAME."""
    text = (html or "").lower()
    markers = [
        ("greenhouse", ("greenhouse.io", "greenhouse-job-board", "gh_jid")),
        ("lever", ("lever.co", "lever-jobs", "lever-postings")),
        ("ashby", ("ashbyhq.com", "ashby-job-board")),
        ("workday", ("myworkdayjobs.com", "/wday/cxs/", "workdayjobs")),
        ("smartrecruiters", ("smartrecruiters.com", "smartrecruiters")),
        ("teamtailor", ("teamtailor", "teamtailor-cdn")),
        ("recruitee", ("recruitee.com", "recruitee-careers")),
        ("personio", ("personio.de", "personio.com", "personio-job")),
        ("icims", ("icims.com", "icimscloud")),
        ("jobvite", ("jobvite.com", "jv-careers")),
        ("breezyhr", ("breezy.hr", "breezy-hr")),
    ]
    for platform, keys in markers:
        if any(k in text for k in keys):
            return platform
    return detect_platform_from_url(base_url)


def extract_known_ats_urls(html: str, base_url: str) -> list[str]:
    soup = BeautifulSoup(html or "", "html.parser")
    urls: list[str] = []
    for tag in soup.find_all(["a", "iframe", "script"]):
        for attr in ("href", "src"):
            raw = tag.get(attr)
            if not raw:
                continue
            u = urljoin(base_url, raw)
            if detect_platform_from_url(u) or "teamtailor" in host_of(u):
                urls.append(u)
    # URLs incrustadas en JSON/JS.
    for raw in re.findall(r'https?:\\?/\\?/[A-Za-z0-9._~:/?#[\]@!$&\'()*+,;=%-]+', html or ""):
        u = raw.replace("\\/", "/")
        if detect_platform_from_url(u) or "teamtailor" in host_of(u):
            urls.append(u)
    out, seen = [], set()
    for u in urls:
        c = canonicalize_url(u)
        if c and c not in seen:
            out.append(u)
            seen.add(c)
    return out


def choose_ats_url(urls: list[str], hinted: str = "") -> str:
    hinted = KNOWN_PLATFORM_ALIASES.get(hinted, hinted)
    if hinted and hinted != "unknown":
        for u in urls:
            if detect_platform_from_url(u) == hinted or (hinted == "teamtailor" and "teamtailor" in host_of(u)):
                return u
    priority = ["greenhouse", "lever", "ashby", "workable", "workday", "smartrecruiters", "recruitee", "personio", "icims", "jobvite", "breezyhr"]
    for p in priority:
        for u in urls:
            if detect_platform_from_url(u) == p:
                return u
    return urls[0] if urls else ""


def extract_greenhouse_slug(url: str) -> str:
    p = urlparse(url)
    qs = parse_qs(p.query)
    if qs.get("for"):
        return qs["for"][0]
    parts = [x for x in p.path.split("/") if x]
    # /embed/job_board?for=company no debe devolver "embed".
    if parts[:2] == ["embed", "job_board"]:
        return ""
    return parts[0] if parts else ""


def fetch_greenhouse(session, company, url) -> list[Job]:
    slug = extract_greenhouse_slug(url)
    if not slug:
        raise RuntimeError("Greenhouse sin slug")
    eu = ".eu.greenhouse.io" in host_of(url)
    api = f"https://boards-api.{ 'eu.' if eu else '' }greenhouse.io/v1/boards/{slug}/jobs?content=true"
    data = session_get(session, api).json()
    jobs = []
    for x in data.get("jobs", []):
        jobs.append(Job(company, x.get("title", ""), x.get("absolute_url", ""), "greenhouse",
                        clean_text((x.get("location") or {}).get("name", "")), clean_text(x.get("content", "")), str(x.get("id", ""))))
    return jobs


def fetch_lever(session, company, url) -> list[Job]:
    parts = [p for p in urlparse(url).path.split("/") if p]
    slug = parts[0] if parts else ""
    if not slug:
        raise RuntimeError("Lever sin slug")
    data = session_get(session, f"https://api.lever.co/v0/postings/{slug}?mode=json").json()
    out = []
    for x in data if isinstance(data, list) else []:
        cats = x.get("categories") or {}
        loc = cats.get("location") or ", ".join(cats.get("allLocations") or [])
        desc = " ".join([x.get("descriptionPlain", ""), x.get("additionalPlain", "")])
        out.append(Job(company, x.get("text", ""), x.get("hostedUrl", ""), "lever", clean_text(loc), clean_text(desc), str(x.get("id", ""))))
    return out


def fetch_ashby(session, company, url) -> list[Job]:
    parts = [p for p in urlparse(url).path.split("/") if p]
    slug = parts[0] if parts else ""
    if not slug:
        raise RuntimeError("Ashby sin slug")
    data = session_get(session, f"https://api.ashbyhq.com/posting-api/job-board/{slug}?includeCompensation=true").json()
    out = []
    for x in data.get("jobs", []) if isinstance(data, dict) else []:
        if not x.get("isListed", True):
            continue
        loc = x.get("location") or ""
        if isinstance(loc, dict):
            loc = loc.get("name", "")
        desc = x.get("descriptionHtml") or x.get("descriptionPlain") or ""
        out.append(Job(company, x.get("title", ""), x.get("jobUrl") or x.get("applyUrl") or "", "ashby", clean_text(str(loc)), clean_text(desc), str(x.get("id", ""))))
    return out


def workable_slug(url: str) -> str:
    p = urlparse(url)
    parts = [x for x in p.path.split("/") if x]
    if "apply.workable.com" in p.netloc and parts:
        return parts[0]
    if p.netloc.endswith(".workable.com") and p.netloc != "www.workable.com":
        return p.netloc.split(".")[0]
    return ""


def fetch_workable(session, company, url) -> list[Job]:
    slug = workable_slug(url)
    if not slug:
        raise RuntimeError("Workable sin account slug")
    data = session_get(session, f"https://www.workable.com/api/accounts/{slug}?details=true").json()
    out = []
    for x in data.get("jobs", []):
        loc = ", ".join([str(v) for v in [x.get("city"), x.get("state"), x.get("country")] if v])
        out.append(Job(company, x.get("title", ""), x.get("application_url") or x.get("url") or x.get("shortlink") or "", "workable", loc, clean_text(x.get("description", "")), str(x.get("shortcode") or x.get("code") or "")))
    return out


def parse_workday(url: str) -> tuple[str, str, str, str]:
    p = urlparse(url)
    host = p.netloc
    if "myworkdayjobs.com" not in host and "myworkdaysite.com" not in host:
        raise RuntimeError("URL no es Workday")
    tenant = host.split(".")[0]
    parts = [x for x in p.path.split("/") if x]
    locale = "en-US"
    if parts and re.fullmatch(r"[a-z]{2}-[A-Z]{2}", parts[0]):
        locale = parts.pop(0)
    site = parts[0] if parts else "External"
    return host, tenant, site, locale


def fetch_workday(session, company, url) -> list[Job]:
    host, tenant, site, locale = parse_workday(url)
    endpoint = f"https://{host}/wday/cxs/{tenant}/{site}/jobs"
    out = []
    offset = 0
    total = None
    while True:
        payload = {"appliedFacets": {}, "limit": 20, "offset": offset, "searchText": ""}
        data = session_post(session, endpoint, json=payload, headers={"Content-Type": "application/json", "Referer": url}).json()
        page = data.get("jobPostings", [])
        if total is None:
            total = data.get("total", len(page))
        if not page:
            break
        for x in page:
            ext = x.get("externalPath", "")
            job_url = f"https://{host}/{locale}/{site}{ext}" if ext else url
            job_id = (x.get("bulletFields") or [""])[-1] if x.get("bulletFields") else ext.rsplit("_", 1)[-1]
            out.append(Job(company, x.get("title", ""), job_url, "workday", clean_text(x.get("locationsText", "")), "", str(job_id)))
        offset += len(page)
        if offset >= (total or 0) or len(page) < 20:
            break
        if offset > 3000:
            break
        time.sleep(REQUEST_DELAY)
    # Solo pedimos descripción para ofertas que potencialmente interesan.
    for j in out:
        if is_design_title(j.title):
            try:
                path = urlparse(j.url).path
                marker = f"/{locale}/{site}"
                external_path = path.split(marker, 1)[1] if marker in path else ""
                detail_url = f"https://{host}/wday/cxs/{tenant}/{site}/job{external_path}"
                info = session_get(session, detail_url, headers={"Referer": url}).json().get("jobPostingInfo", {})
                j.description = clean_text(info.get("jobDescription", ""))
                j.location = j.location or clean_text(info.get("location", ""))
            except Exception:
                pass
    return out


def fetch_smartrecruiters(session, company, url) -> list[Job]:
    parts = [p for p in urlparse(url).path.split("/") if p]
    slug = parts[0] if parts else ""
    if not slug:
        raise RuntimeError("SmartRecruiters sin company slug")
    data = session_get(session, f"https://api.smartrecruiters.com/v1/companies/{slug}/postings").json()
    out = []
    for x in data.get("content", []):
        ref = (x.get("ref") or {}).get("jobAd", {})
        job_url = ref.get("sourceUrl") or f"https://careers.smartrecruiters.com/{slug}/{x.get('id', '')}"
        loc = x.get("location") or {}
        loc_text = ", ".join(str(loc.get(k, "")) for k in ("city", "region", "country") if loc.get(k)) if isinstance(loc, dict) else str(loc)
        out.append(Job(company, x.get("name", ""), job_url, "smartrecruiters", loc_text, "", str(x.get("id", ""))))
    return out


def fetch_recruitee(session, company, url) -> list[Job]:
    origin = f"{urlparse(url).scheme or 'https'}://{urlparse(url).netloc}"
    data = session_get(session, f"{origin}/api/offers/").json()
    offers = data.get("offers", []) if isinstance(data, dict) else []
    out = []
    for x in offers:
        locs = x.get("locations") or []
        if isinstance(locs, list):
            loc = ", ".join(clean_text(str(v.get("name", "") if isinstance(v, dict) else v)) for v in locs)
        else:
            loc = clean_text(str(locs))
        out.append(Job(company, x.get("title", ""), x.get("careers_url") or x.get("url") or "", "recruitee", loc, clean_text(x.get("description", "")), str(x.get("id", ""))))
    return out


def fetch_personio(session, company, url) -> list[Job]:
    p = urlparse(url)
    host = p.netloc
    xml_url = f"https://{host}/xml?language=en"
    r = session_get(session, xml_url)
    root = ET.fromstring(r.text)
    out = []
    for pos in root.findall(".//position"):
        title = (pos.findtext("name") or "").strip()
        pid = (pos.findtext("id") or "").strip()
        office = (pos.findtext("office") or "").strip()
        desc_parts = []
        for val in pos.findall(".//jobDescription/value"):
            if val.text:
                desc_parts.append(val.text)
        job_url = f"https://{host}/job/{pid}" if pid else url
        out.append(Job(company, title, job_url, "personio", office, clean_text(" ".join(desc_parts)), pid))
    return out


def generic_jobs_from_html(company: str, html: str, base_url: str, platform: str = "html") -> list[Job]:
    soup = BeautifulSoup(html or "", "html.parser")
    out: list[Job] = []

    # JSON-LD JobPosting: es la señal más fiable en una web corporativa.
    for tag in soup.find_all("script", attrs={"type": re.compile("ld\\+json", re.I)}):
        text = tag.string or tag.get_text("", strip=True)
        if not text:
            continue
        try:
            payload = json.loads(text)
        except Exception:
            continue
        nodes = payload if isinstance(payload, list) else [payload]
        expanded = []
        for node in nodes:
            if isinstance(node, dict) and isinstance(node.get("@graph"), list):
                expanded.extend(node["@graph"])
            else:
                expanded.append(node)
        for node in expanded:
            if not isinstance(node, dict):
                continue
            node_type = node.get("@type")
            types = node_type if isinstance(node_type, list) else [node_type]
            if "JobPosting" not in types:
                continue
            title = str(node.get("title") or "")
            job_url = str(node.get("url") or base_url)
            desc = clean_text(str(node.get("description") or ""))
            loc = ""
            jl = node.get("jobLocation")
            if isinstance(jl, dict):
                jl = [jl]
            if isinstance(jl, list):
                bits = []
                for item in jl:
                    addr = (item or {}).get("address", {}) if isinstance(item, dict) else {}
                    if isinstance(addr, dict):
                        bits.append(", ".join(str(addr.get(k, "")) for k in ("addressLocality", "addressRegion", "addressCountry") if addr.get(k)))
                loc = "; ".join(x for x in bits if x)
            ext = str(node.get("identifier") or "")
            out.append(Job(company, title, job_url, platform, loc, desc, ext))

    # Enlaces de vacantes. Incluimos títulos de diseño y también enlaces claramente
    # "job-like" para poder afirmar con más confianza que no había puestos de diseño.
    generic_labels = {"jobs", "job", "careers", "career", "view jobs", "open roles", "open positions", "apply", "learn more", "see all jobs"}
    for a in soup.find_all("a", href=True):
        title = clean_text(a.get_text(" ", strip=True) or a.get("title", ""))
        href = urljoin(base_url, a["href"])
        if not title or is_aggregator(href):
            continue
        path = urlparse(href).path.lower()
        jobish = bool(re.search(r"/(?:jobs?|positions?|vacancies|openings?)/[^/]+", path))
        if is_design_title(title) or (jobish and norm(title) not in generic_labels and len(title) >= 4):
            out.append(Job(company, title, href, platform))

    # Deduplicar.
    dedup = {}
    for j in out:
        key = (norm(j.title), j.canonical_url)
        if j.title and j.url:
            dedup[key] = j
    return list(dedup.values())


def candidate_listing_urls(html: str, base_url: str) -> list[str]:
    """Encuentra páginas internas tipo Open roles / Vacancies / Jobs (un solo salto)."""
    soup = BeautifulSoup(html or "", "html.parser")
    base_host = host_of(base_url)
    scored = []
    tokens = ("career", "jobs", "job", "vacan", "open role", "open position", "positions", "opportun", "join us", "join-us")
    for a in soup.find_all("a", href=True):
        href = urljoin(base_url, a.get("href") or "")
        if not href.startswith("http") or is_aggregator(href):
            continue
        if host_of(href) != base_host and not detect_platform_from_url(href):
            continue
        label = norm(a.get_text(" ", strip=True))
        path = norm(urlparse(href).path.replace("-", " ").replace("_", " "))
        score = sum(2 for t in tokens if t in label) + sum(1 for t in tokens if t in path)
        if detect_platform_from_url(href):
            score += 5
        if score:
            scored.append((score, href))
    seen, out = set(), []
    for _, href in sorted(scored, key=lambda x: x[0], reverse=True):
        c = canonicalize_url(href)
        if c in seen or c == canonicalize_url(base_url):
            continue
        seen.add(c)
        out.append(href)
        if len(out) >= MAX_GENERIC_PAGES:
            break
    return out


def fetch_generic_portal(session, company: str, url: str, browser: BrowserFetcher, label: str = "generic", avoid_platform: str = "") -> tuple[list[Job], str]:
    """Fallback robusto para portales propios o APIs que han fallado.

    Lee landing + hasta unas pocas páginas internas y, si hace falta, el DOM renderizado.
    Puede saltar a otro ATS detectado, pero evita volver al adaptador que acaba de fallar.
    """
    r = session_get(session, url)
    pages = [(r.text, r.url, "html")]
    for candidate in candidate_listing_urls(r.text, r.url):
        p = detect_platform_from_url(candidate)
        if p and p != avoid_platform:
            try:
                dummy = CompanyRow(0, company, candidate, p, "")
                jobs, strat = fetch_by_platform(session, dummy, candidate, p, browser)
                return jobs, f"{label}_redirect_{strat}"
            except Exception:
                pass
        try:
            rr = session_get(session, candidate)
            pages.append((rr.text, rr.url, "html"))
        except Exception:
            continue

    jobs = []
    seen = set()
    for page_html, page_url, _ in pages:
        for j in generic_jobs_from_html(company, page_html, page_url, label):
            key = (norm(j.title), j.canonical_url)
            if key not in seen:
                jobs.append(j); seen.add(key)
    if jobs:
        return jobs, f"{label}_html_crawl"

    if browser.enabled:
        browser_targets = [x[1] for x in pages[:3]]
        for target in browser_targets:
            try:
                dyn_html, final_url = browser.fetch(target)
                ats_urls = extract_known_ats_urls(dyn_html, final_url)
                chosen = choose_ats_url(ats_urls)
                p = detect_platform_from_url(chosen) if chosen else ""
                if chosen and p and p != avoid_platform:
                    dummy = CompanyRow(0, company, chosen, p, "")
                    alt, strat = fetch_by_platform(session, dummy, chosen, p, browser)
                    return alt, f"{label}_browser_redirect_{strat}"
                for j in generic_jobs_from_html(company, dyn_html, final_url, label):
                    key = (norm(j.title), j.canonical_url)
                    if key not in seen:
                        jobs.append(j); seen.add(key)
            except Exception:
                continue
        return jobs, f"{label}_browser_crawl"
    return [], f"{label}_html_crawl"


def fetch_teamtailor_or_generic(session, company, url, browser: BrowserFetcher) -> tuple[list[Job], str]:
    # Los career sites Teamtailor suelen exponer la lista bajo /jobs y son perfectamente
    # extraíbles desde el HTML renderizado, aunque no exista un feed público uniforme.
    try_urls = [url]
    p = urlparse(url)
    origin = f"{p.scheme or 'https'}://{p.netloc}"
    if not p.path.rstrip("/").endswith("/jobs"):
        try_urls.append(origin + "/jobs")
    last_error = ""
    for candidate in try_urls:
        try:
            r = session_get(session, candidate)
            jobs = generic_jobs_from_html(company, r.text, r.url, "teamtailor")
            if jobs:
                return jobs, "teamtailor_html"
            ats_urls = extract_known_ats_urls(r.text, r.url)
            if ats_urls:
                return [], f"redirect:{choose_ats_url(ats_urls)}"
            if browser.enabled:
                html, final_url = browser.fetch(r.url)
                jobs = generic_jobs_from_html(company, html, final_url, "teamtailor")
                if jobs:
                    return jobs, "teamtailor_browser"
        except Exception as exc:
            last_error = str(exc)
    if last_error:
        raise RuntimeError(last_error)
    return [], "teamtailor_html"


def discover_official_careers(session: requests.Session, company: str) -> str:
    q = quote_plus(f'"{company}" careers jobs')
    url = f"https://html.duckduckgo.com/html/?q={q}"
    r = session_get(session, url, headers={"Referer": "https://duckduckgo.com/"})
    soup = BeautifulSoup(r.text, "html.parser")
    candidates = []
    company_tokens = [t for t in re.findall(r"[a-z0-9]+", norm(company)) if len(t) >= 4][:4]
    for a in soup.select("a.result__a, .results_links a[href], a[href]"):
        href = a.get("href") or ""
        text = clean_text(a.get_text(" ", strip=True))
        if not href:
            continue
        # DDG envuelve algunos resultados en /l/?uddg=...
        if "duckduckgo.com/l/" in href or href.startswith("/l/"):
            qs = parse_qs(urlparse(urljoin("https://duckduckgo.com", href)).query)
            href = (qs.get("uddg") or [href])[0]
        if not href.startswith("http") or is_aggregator(href):
            continue
        h = host_of(href)
        path = urlparse(href).path.lower()
        score = 0
        if detect_platform_from_url(href):
            score += 6
        if any(k in path for k in ("career", "jobs", "join", "vacan", "work-with-us", "opportun")):
            score += 4
        if any(t in h.replace("-", "") or t in norm(text).replace(" ", "") for t in company_tokens):
            score += 4
        if any(k in norm(text) for k in ("careers", "jobs", "join us", "vacancies")):
            score += 2
        candidates.append((score, href))
    candidates.sort(key=lambda x: x[0], reverse=True)
    for score, href in candidates[:8]:
        if score < 4:
            continue
        try:
            resp = session_get(session, href)
            if not is_aggregator(resp.url):
                return resp.url
        except Exception:
            continue
    return ""


def resolve_company(session: requests.Session, row: CompanyRow, cache: dict, browser: BrowserFetcher) -> tuple[str, str, str]:
    key = norm(row.company)
    cached = cache.get(key, {})
    hint = KNOWN_PLATFORM_ALIASES.get(row.platform_hint, row.platform_hint)

    input_url = row.input_url
    if (not input_url or is_aggregator(input_url)) and cached.get("resolved_url") and cache_is_fresh(cached):
        return cached["resolved_url"], cached.get("platform", ""), "cache"

    url = input_url
    source = "input"
    if not url or is_aggregator(url):
        url = discover_official_careers(session, row.company)
        source = "search"
        if not url:
            return "", "", "unresolved_no_official_careers_url"

    direct_platform = detect_platform_from_url(url)
    if direct_platform:
        return url, direct_platform, source
    if hint == "teamtailor":
        return url, "teamtailor", source

    # Algunas pistas de ATS solo sirven si tenemos la URL real del board. Una URL
    # corporativa marcada como "workday"/"greenhouse" no debe enviarse directamente
    # al adaptador: primero hay que descubrir el enlace ATS real.
    direct_url_required = {"greenhouse", "lever", "ashby", "workable", "workday", "smartrecruiters", "personio"}

    # Para pistas como Workday/iCIMS con URL corporativa, inspeccionamos el portal para
    # localizar el enlace ATS real; lo mismo para unknown/otro/verificar.
    try:
        r = session_get(session, url)
        final_url = r.url
        p = detect_platform_from_url(final_url)
        if p:
            return final_url, p, source + "+redirect"
        html_platform = detect_platform_from_html(r.text, final_url)
        ats_urls = extract_known_ats_urls(r.text, final_url)
        chosen = choose_ats_url(ats_urls, hint or html_platform)
        if chosen:
            return chosen, detect_platform_from_url(chosen) or hint or "unknown", source + "+html_detect"
        if browser.enabled and hint in direct_url_required:
            try:
                dyn_html, dyn_final = browser.fetch(final_url)
                dyn_urls = extract_known_ats_urls(dyn_html, dyn_final)
                dyn_chosen = choose_ats_url(dyn_urls, hint)
                if dyn_chosen:
                    return dyn_chosen, detect_platform_from_url(dyn_chosen) or hint, source + "+browser_detect"
            except Exception:
                pass
        # Mantener la web oficial como estrategia genérica. Para ATS que requieren URL
        # directa, degradamos a generic si no pudimos encontrar el board real.
        fallback_platform = html_platform or (hint if hint not in ("", "unknown") and hint not in direct_url_required else "generic")
        return final_url, fallback_platform, source + "+generic"
    except Exception:
        # Browser puede rescatar portales que bloquean requests.
        if browser.enabled:
            try:
                html, final_url = browser.fetch(url)
                p = detect_platform_from_url(final_url)
                if p:
                    return final_url, p, source + "+browser_redirect"
                html_platform = detect_platform_from_html(html, final_url)
                ats_urls = extract_known_ats_urls(html, final_url)
                chosen = choose_ats_url(ats_urls, hint or html_platform)
                if chosen:
                    return chosen, detect_platform_from_url(chosen) or hint or "unknown", source + "+browser_detect"
                return final_url, html_platform or (hint if hint not in ("", "unknown") else "generic"), source + "+browser_generic"
            except Exception:
                pass
        fallback_platform = hint if hint not in ("", "unknown") and hint not in direct_url_required else "generic"
        return url, fallback_platform, source + "+fetch_failed"


def fetch_by_platform(session, row: CompanyRow, resolved_url: str, platform: str, browser: BrowserFetcher) -> tuple[list[Job], str]:
    if platform == "greenhouse":
        return fetch_greenhouse(session, row.company, resolved_url), "greenhouse_api"
    if platform == "lever":
        return fetch_lever(session, row.company, resolved_url), "lever_api"
    if platform == "ashby":
        return fetch_ashby(session, row.company, resolved_url), "ashby_api"
    if platform == "workable":
        return fetch_workable(session, row.company, resolved_url), "workable_api"
    if platform == "workday":
        return fetch_workday(session, row.company, resolved_url), "workday_cxs"
    if platform == "smartrecruiters":
        return fetch_smartrecruiters(session, row.company, resolved_url), "smartrecruiters_api"
    if platform == "recruitee":
        return fetch_recruitee(session, row.company, resolved_url), "recruitee_api"
    if platform == "personio":
        return fetch_personio(session, row.company, resolved_url), "personio_xml"
    if platform == "teamtailor":
        jobs, strategy = fetch_teamtailor_or_generic(session, row.company, resolved_url, browser)
        if strategy.startswith("redirect:"):
            redirect = strategy.split(":", 1)[1]
            p = detect_platform_from_url(redirect)
            if p and p != "teamtailor":
                return fetch_by_platform(session, row, redirect, p, browser)
        return jobs, strategy

    # iCIMS, Jobvite, BreezyHR y portales propios: landing + páginas internas + browser.
    return fetch_generic_portal(session, row.company, resolved_url, browser, platform or "generic")


def fetch_with_recovery(session, row: CompanyRow, resolved_url: str, platform: str, browser: BrowserFetcher) -> tuple[list[Job], str, str]:
    """Intenta el adaptador principal y degrada a portal genérico si la API cambia/falla."""
    try:
        jobs, strategy = fetch_by_platform(session, row, resolved_url, platform, browser)
        return jobs, strategy, ""
    except Exception as primary:
        primary_msg = f"{type(primary).__name__}: {primary}"
        try:
            jobs, strategy = fetch_generic_portal(
                session, row.company, resolved_url, browser,
                label=f"{platform or 'generic'}_fallback", avoid_platform=platform,
            )
            return jobs, strategy, f"Adaptador principal falló ({primary_msg}); usado fallback web."
        except Exception:
            raise primary


def write_csv(path: Path, rows: list[dict], headers: list[str]):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def render_report(coverage: list[dict], new_jobs: list[Job], open_jobs: list[Job]) -> str:
    counts = {}
    for r in coverage:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    ok = sum(v for k, v in counts.items() if k.startswith("checked"))
    unresolved = sum(v for k, v in counts.items() if k.startswith("unresolved"))
    errors = sum(v for k, v in counts.items() if k == "error")
    lines = [
        f"# Búsqueda automática — {datetime.now().strftime('%d-%m-%Y %H:%M')}",
        "",
        f"- Empresas/fila procesadas: **{len(coverage)}**",
        f"- Revisadas con éxito: **{ok}**",
        f"- Sin resolver todavía: **{unresolved}**",
        f"- Errores: **{errors}**",
        f"- Ofertas de diseño abiertas detectadas: **{len(open_jobs)}**",
        f"- Ofertas nuevas: **{len(new_jobs)}**",
        "",
        "## Ofertas nuevas",
        "",
    ]
    if new_jobs:
        for j in new_jobs:
            loc = f" — {j.location}" if j.location else ""
            lines.append(f"- **{j.company} — {j.title}**{loc}: {j.url}")
    else:
        lines.append("No se detectaron ofertas nuevas de diseño.")
    lines += ["", "## Cobertura", ""]
    for status, count in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
        lines.append(f"- {status}: {count}")
    bad = [r for r in coverage if r["status"] in ("error", "unresolved") or r["status"].startswith("unresolved")]
    if bad:
        lines += ["", "## Pendientes / errores", ""]
        for r in bad:
            detail = r.get("error") or r.get("warning") or r.get("resolution_source") or ""
            lines.append(f"- **{r['empresa']}** — {r['status']}: {detail}")
    return "\n".join(lines) + "\n"


def render_dashboard(coverage: list[dict], new_jobs: list[Job], open_jobs: list[Job]) -> str:
    total = len(coverage)
    checked = sum(1 for r in coverage if str(r.get("status", "")).startswith("checked"))
    errors = sum(1 for r in coverage if r.get("status") == "error")
    unresolved = sum(1 for r in coverage if str(r.get("status", "")).startswith("unresolved"))
    pct = round(100 * checked / total, 1) if total else 0

    def e(v):
        return html_lib.escape(str(v or ""), quote=True)

    def job_cards():
        if not new_jobs:
            return '<div class="empty">No se detectaron ofertas nuevas en esta ejecución.</div>'
        cards = []
        for j in new_jobs:
            loc = f'<span>{e(j.location)}</span>' if j.location else ''
            cards.append(f'''<article class="job-card">
              <div class="job-meta"><span>{e(j.company)}</span>{loc}<span>{e(j.platform)}</span></div>
              <h3>{e(j.title)}</h3>
              <a href="{e(j.url)}" target="_blank" rel="noopener">Abrir oferta →</a>
            </article>''')
        return "".join(cards)

    bad = [r for r in coverage if r.get("status") == "error" or str(r.get("status", "")).startswith("unresolved")]
    def bad_rows():
        if not bad:
            return '<tr><td colspan="5" class="empty-cell">Sin pendientes ni errores.</td></tr>'
        rows = []
        for r in bad:
            detail = r.get("error") or r.get("warning") or r.get("resolution_source") or ""
            url = r.get("resolved_url") or r.get("input_url") or ""
            link = f'<a href="{e(url)}" target="_blank" rel="noopener">Abrir</a>' if str(url).startswith("http") else "—"
            rows.append(f'''<tr><td><strong>{e(r.get("empresa"))}</strong></td><td><span class="badge bad">{e(r.get("status"))}</span></td><td>{e(r.get("detected_platform")) or "—"}</td><td>{e(detail)}</td><td>{link}</td></tr>''')
        return "".join(rows)

    all_rows = []
    for r in coverage:
        st = str(r.get("status", ""))
        cls = "ok" if st.startswith("checked") else "bad" if st == "error" else "warn"
        url = r.get("resolved_url") or r.get("input_url") or ""
        link = f'<a href="{e(url)}" target="_blank" rel="noopener">Abrir</a>' if str(url).startswith("http") else "—"
        all_rows.append(f'''<tr data-status="{e(st)}"><td>{e(r.get("empresa"))}</td><td><span class="badge {cls}">{e(st)}</span></td><td>{e(r.get("detected_platform")) or "—"}</td><td>{e(r.get("strategy")) or "—"}</td><td>{e(r.get("design_jobs_found"))}</td><td>{e(r.get("seconds"))}s</td><td>{link}</td></tr>''')

    generated = datetime.now().strftime('%d/%m/%Y · %H:%M')
    return f'''<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Job Watch · {e(generated)}</title>
<style>
:root{{--bg:#f6f7f9;--panel:#fff;--ink:#18202a;--muted:#667085;--line:#e3e7ed;--ok:#166534;--okbg:#dcfce7;--warn:#854d0e;--warnbg:#fef9c3;--bad:#991b1b;--badbg:#fee2e2;--accent:#294e8c}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 system-ui,-apple-system,Segoe UI,sans-serif}}a{{color:var(--accent);font-weight:650;text-decoration:none}}a:hover{{text-decoration:underline}}.wrap{{max-width:1180px;margin:auto;padding:32px 20px 70px}}header{{display:flex;justify-content:space-between;gap:20px;align-items:end;margin-bottom:22px}}h1{{margin:0;font-size:30px;letter-spacing:-.03em}}.sub{{color:var(--muted);margin-top:5px}}.grid{{display:grid;grid-template-columns:repeat(5,1fr);gap:12px;margin:18px 0 28px}}.stat,.panel,.job-card{{background:var(--panel);border:1px solid var(--line);border-radius:14px}}.stat{{padding:16px}}.stat b{{display:block;font-size:25px}}.stat span{{color:var(--muted);font-size:13px}}.panel{{padding:20px;margin:18px 0}}h2{{font-size:19px;margin:0 0 14px}}.jobs{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px}}.job-card{{padding:17px}}.job-card h3{{margin:7px 0 12px;font-size:17px}}.job-meta{{display:flex;gap:8px;flex-wrap:wrap;color:var(--muted);font-size:12px}}.job-meta span:not(:last-child)::after{{content:' ·';margin-left:8px}}.badge{{display:inline-block;border-radius:999px;padding:3px 8px;font-size:12px;font-weight:700;white-space:nowrap}}.badge.ok{{background:var(--okbg);color:var(--ok)}}.badge.warn{{background:var(--warnbg);color:var(--warn)}}.badge.bad{{background:var(--badbg);color:var(--bad)}}.table-wrap{{overflow:auto;border:1px solid var(--line);border-radius:10px}}table{{width:100%;border-collapse:collapse;background:#fff}}th,td{{text-align:left;padding:10px 12px;border-bottom:1px solid var(--line);vertical-align:top}}th{{font-size:12px;color:var(--muted);background:#fafbfc;position:sticky;top:0}}tr:last-child td{{border-bottom:0}}.empty,.empty-cell{{color:var(--muted);padding:18px}}details>summary{{cursor:pointer;font-weight:700;margin-bottom:14px}}.bar{{height:8px;border-radius:9px;background:#e8ebef;overflow:hidden;margin-top:9px}}.bar>i{{display:block;height:100%;width:{pct}%;background:#2f855a}}.note{{color:var(--muted);font-size:13px}}@media(max-width:800px){{.grid{{grid-template-columns:repeat(2,1fr)}}.jobs{{grid-template-columns:1fr}}header{{display:block}}}}
</style></head><body><main class="wrap">
<header><div><h1>Job Watch</h1><div class="sub">Última ejecución: {e(generated)}</div></div><div class="note">La cobertura cuenta empresas realmente revisadas, no solo filas leídas.</div></header>
<section class="grid">
<div class="stat"><b>{total}</b><span>empresas procesadas</span></div><div class="stat"><b>{checked}</b><span>revisadas</span></div><div class="stat"><b>{pct}%</b><span>cobertura real</span><div class="bar"><i></i></div></div><div class="stat"><b>{len(new_jobs)}</b><span>ofertas nuevas</span></div><div class="stat"><b>{errors + unresolved}</b><span>pendientes / errores</span></div>
</section>
<section class="panel"><h2>Ofertas nuevas</h2><div class="jobs">{job_cards()}</div></section>
<section class="panel"><h2>Lo que necesita atención</h2><div class="table-wrap"><table><thead><tr><th>Empresa</th><th>Estado</th><th>ATS</th><th>Detalle</th><th>Link</th></tr></thead><tbody>{bad_rows()}</tbody></table></div></section>
<section class="panel"><details><summary>Ver cobertura completa ({total} empresas · {len(open_jobs)} ofertas de diseño abiertas)</summary><div class="table-wrap"><table><thead><tr><th>Empresa</th><th>Estado</th><th>ATS</th><th>Método</th><th>Diseño</th><th>Tiempo</th><th>Link</th></tr></thead><tbody>{''.join(all_rows)}</tbody></table></div></details></section>
</main></body></html>'''


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--watchlist", default=str(WATCHLIST))
    ap.add_argument("--browser", action="store_true", help="Usar Playwright como fallback para portales JS")
    ap.add_argument("--company", default="", help="Procesar solo empresas cuyo nombre contenga este texto")
    ap.add_argument("--max-companies", type=int, default=0)
    args = ap.parse_args()

    path = Path(args.watchlist)
    if not path.exists():
        raise SystemExit(f"No existe {path}")

    rows = load_watchlist(path)
    if args.company:
        needle = norm(args.company)
        rows = [r for r in rows if needle in norm(r.company)]
    if args.max_companies:
        rows = rows[: args.max_companies]

    session = build_session()
    browser = BrowserFetcher(args.browser)
    cache = load_cache()
    seen_urls, seen_signatures, seen_fingerprints = load_seen()

    coverage: list[dict] = []
    open_design_jobs: list[Job] = []
    new_jobs: list[Job] = []
    run_seen: set[str] = set()
    run_fingerprints: set[str] = set()

    try:
        for i, row in enumerate(rows, 1):
            print(f"[{i}/{len(rows)}] {row.company}")
            started = time.time()
            resolved_url = ""
            detected = ""
            resolution_source = ""
            strategy = ""
            error = ""
            warning = ""
            all_jobs: list[Job] = []
            try:
                resolved_url, detected, resolution_source = resolve_company(session, row, cache, browser)
                if not resolved_url:
                    status = "unresolved"
                else:
                    try:
                        all_jobs, strategy, warning = fetch_with_recovery(session, row, resolved_url, detected, browser)
                    except Exception:
                        # Si una resolución cacheada envejeció o cambió de ATS, descartarla
                        # y redescubrir una sola vez antes de marcar error.
                        if resolution_source == "cache":
                            cache.pop(norm(row.company), None)
                            resolved_url, detected, resolution_source = resolve_company(session, row, cache, browser)
                            if not resolved_url:
                                raise
                            all_jobs, strategy, warning = fetch_with_recovery(session, row, resolved_url, detected, browser)
                        else:
                            raise
                    design_jobs = [j for j in all_jobs if is_design_title(j.title) and j.url and not is_aggregator(j.url)]
                    # Dedup por firma dentro de la misma ejecución.
                    unique = []
                    for j in design_jobs:
                        fp = j.content_fingerprint
                        if j.signature in run_seen or (fp and fp in run_fingerprints):
                            continue
                        run_seen.add(j.signature)
                        if fp:
                            run_fingerprints.add(fp)
                        unique.append(j)
                    design_jobs = unique
                    open_design_jobs.extend(design_jobs)
                    for j in design_jobs:
                        fp = j.content_fingerprint
                        already_seen = (
                            j.canonical_url in seen_urls or
                            j.signature in seen_signatures or
                            (fp and fp in seen_fingerprints)
                        )
                        if not already_seen:
                            new_jobs.append(j)
                            seen_urls.add(j.canonical_url)
                            seen_signatures.add(j.signature)
                            if fp:
                                seen_fingerprints.add(fp)
                    if design_jobs:
                        status = "checked_design_found"
                    elif not all_jobs and (
                        "generic" in strategy or "fallback" in strategy or
                        detected in {"icims", "jobvite", "breezyhr", "unknown", "generic"}
                    ):
                        # Una web propia que no nos deja identificar ni una sola vacante
                        # no debe contarse como "sin diseño": sería un falso negativo.
                        status = "unresolved_listing_not_parsed"
                    else:
                        status = "checked_no_design"
                    cache[norm(row.company)] = {
                        "company": row.company,
                        "resolved_url": resolved_url,
                        "platform": detected,
                        "last_success": datetime.now().isoformat(timespec="seconds"),
                        "strategy": strategy,
                    }
            except Exception as exc:
                status = "error"
                error = f"{type(exc).__name__}: {exc}"
                print(f"  ERROR {error}", file=sys.stderr)

            coverage.append({
                "row": row.row_number,
                "empresa": row.company,
                "input_url": row.input_url,
                "platform_hint": row.platform_hint,
                "resolved_url": resolved_url,
                "detected_platform": detected,
                "resolution_source": resolution_source,
                "strategy": strategy,
                "status": status,
                "all_jobs_seen": len(all_jobs),
                "design_jobs_found": len([j for j in all_jobs if is_design_title(j.title)]),
                "warning": warning,
                "error": error,
                "seconds": round(time.time() - started, 2),
                "checked_at": datetime.now().isoformat(timespec="seconds"),
            })
            time.sleep(REQUEST_DELAY)
    finally:
        browser.close()

    append_seen(new_jobs)
    save_cache(cache)

    job_headers = ["empresa", "titulo", "url", "plataforma", "ubicacion", "descripcion", "external_id"]
    def job_row(j: Job):
        return {"empresa": j.company, "titulo": j.title, "url": j.url, "plataforma": j.platform,
                "ubicacion": j.location, "descripcion": j.description, "external_id": j.external_id}
    write_csv(NEW_CSV, [job_row(j) for j in new_jobs], job_headers)
    write_csv(OPEN_CSV, [job_row(j) for j in open_design_jobs], job_headers)
    coverage_headers = ["row", "empresa", "input_url", "platform_hint", "resolved_url", "detected_platform", "resolution_source", "strategy", "status", "all_jobs_seen", "design_jobs_found", "warning", "error", "seconds", "checked_at"]
    write_csv(COVERAGE_CSV, coverage, coverage_headers)
    REPORT_MD.write_text(render_report(coverage, new_jobs, open_design_jobs), encoding="utf-8")
    DASHBOARD_HTML.write_text(render_dashboard(coverage, new_jobs, open_design_jobs), encoding="utf-8")

    checked = sum(1 for r in coverage if r["status"].startswith("checked"))
    unresolved = sum(1 for r in coverage if r["status"].startswith("unresolved"))
    errors = sum(1 for r in coverage if r["status"] == "error")
    print("\n=== RESUMEN ===")
    print(f"Filas procesadas: {len(coverage)}")
    print(f"Revisadas: {checked}")
    print(f"Sin resolver: {unresolved}")
    print(f"Errores: {errors}")
    print(f"Ofertas de diseño abiertas: {len(open_design_jobs)}")
    print(f"Ofertas nuevas: {len(new_jobs)}")
    print(f"Cobertura: {COVERAGE_CSV.name}")
    print(f"Dashboard: {DASHBOARD_HTML.name}")


if __name__ == "__main__":
    main()
