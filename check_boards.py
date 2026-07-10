#!/usr/bin/env python3
"""
check_boards.py

Revisa watchlist_consolidado.xlsx, consulta las APIs públicas de Greenhouse,
Lever y Ashby para las empresas marcadas con esa plataforma, filtra los
puestos de diseño, y descarta lo que ya se ha visto en ejecuciones anteriores.

Escribe las novedades en nuevas_ofertas.csv (se sobrescribe cada ejecución).
Guarda un historial permanente en ofertas_vistas.csv para no repetir ofertas.

Uso:
    python3 check_boards.py

Requiere:
    pip install openpyxl --break-system-packages
"""

import csv
import json
import time
import urllib.error
import urllib.request
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

CARPETA = Path(__file__).parent
WATCHLIST = CARPETA / "watchlist_consolidado.xlsx"
VISTAS_LOG = CARPETA / "ofertas_vistas.csv"
SALIDA = CARPETA / "nuevas_ofertas.csv"

# Palabras que identifican un puesto de diseño de producto / UX-UI.
# Todo en minúsculas, coincidencia por substring sobre el título en minúsculas.
PALABRAS_CLAVE = [
    "product designer", "product design",
    "diseñador de producto", "diseñadora de producto", "diseño de producto",
    "ux/ui", "ux ui", "ui/ux", "ui ux",
    "ux designer", "ui designer",
    "diseñador ux", "diseñadora ux", "diseñador ui", "diseñadora ui",
    "diseñador/a ux", "diseñador/a ui", "diseñador/a de producto",
    "design system", "sistema de diseño", "design systems",
]

USER_AGENT = "Mozilla/5.0 (compatible; busqueda-trabajo-script/1.0)"


def coincide_titulo(titulo):
    if not titulo:
        return False
    t = titulo.lower()
    return any(p in t for p in PALABRAS_CLAVE)


def extraer_slug(url, plataforma):
    """Extrae el 'slug' (token de empresa) de una URL, cuando la URL
    apunta directamente a un dominio de la propia plataforma ATS.
    Si la empresa usa un dominio propio (ej. empresa.com/careers),
    no hay forma de adivinar el slug desde aquí y se devuelve None."""
    try:
        parsed = urlparse(url)
        dominio = parsed.netloc.lower()
        partes = [p for p in parsed.path.split("/") if p]
        if not partes:
            return None
        if plataforma == "greenhouse" and "greenhouse.io" in dominio:
            return partes[0]
        if plataforma == "lever" and "lever.co" in dominio:
            return partes[0]
        if plataforma == "ashby" and "ashbyhq.com" in dominio:
            return partes[0]
    except Exception:
        return None
    return None


def es_greenhouse_eu(url):
    return ".eu.greenhouse.io" in url.lower()


def pedir_json(url):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=15) as r:
        return json.loads(r.read().decode("utf-8"))


def consultar_greenhouse(slug, eu=False):
    dominio = "boards-api.eu.greenhouse.io" if eu else "boards-api.greenhouse.io"
    data = pedir_json(f"https://{dominio}/v1/boards/{slug}/jobs?content=true")
    return [
        {
            "titulo": j.get("title", ""),
            "url": j.get("absolute_url", ""),
        }
        for j in data.get("jobs", [])
    ]


def consultar_lever(slug):
    data = pedir_json(f"https://api.lever.co/v0/postings/{slug}?mode=json")
    return [
        {
            "titulo": j.get("text", ""),
            "url": j.get("hostedUrl", ""),
        }
        for j in data
    ]


def consultar_ashby(slug):
    data = pedir_json(
        f"https://api.ashbyhq.com/posting-api/job-board/{slug}?includeCompensation=false"
    )
    jobs = data.get("jobs", []) if isinstance(data, dict) else []
    return [
        {
            "titulo": j.get("title", ""),
            "url": j.get("jobUrl") or j.get("applyUrl", ""),
        }
        for j in jobs
        if j.get("isListed", True)
    ]


def cargar_vistas():
    if not VISTAS_LOG.exists():
        return set()
    with open(VISTAS_LOG, newline="", encoding="utf-8") as f:
        return {row["url"] for row in csv.DictReader(f) if row.get("url")}


def registrar_vista(writer, empresa, titulo, url):
    writer.writerow(
        {
            "empresa": empresa,
            "titulo": titulo,
            "url": url,
            "fecha_visto": time.strftime("%Y-%m-%d"),
        }
    )


def main():
    if not WATCHLIST.exists():
        print(f"No encuentro {WATCHLIST.name} en esta carpeta. ¿Está el script en busqueda-trabajo/?")
        return

    wb = openpyxl.load_workbook(WATCHLIST, data_only=True)
    ws = wb["watchlist"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    faltan = {"empresa", "web_empleo", "plataforma_ats"} - set(idx)
    if faltan:
        print(f"Faltan columnas en el Excel: {faltan}")
        return

    vistas = cargar_vistas()
    nuevas = []
    sin_slug = []
    errores = []
    revisadas = 0

    vistas_es_nuevo = not VISTAS_LOG.exists()
    vistas_f = open(VISTAS_LOG, "a", newline="", encoding="utf-8")
    vistas_writer = csv.DictWriter(
        vistas_f, fieldnames=["empresa", "titulo", "url", "fecha_visto"]
    )
    if vistas_es_nuevo:
        vistas_writer.writeheader()

    for row in ws.iter_rows(min_row=2, values_only=True):
        empresa = row[idx["empresa"]]
        url_empleo = row[idx["web_empleo"]]
        plataforma = (row[idx["plataforma_ats"]] or "").strip().lower()

        if not empresa or plataforma not in ("greenhouse", "lever", "ashby"):
            continue
        if not url_empleo:
            continue

        slug = extraer_slug(url_empleo, plataforma)
        if not slug:
            sin_slug.append(empresa)
            continue

        revisadas += 1
        try:
            if plataforma == "greenhouse":
                ofertas = consultar_greenhouse(slug, eu=es_greenhouse_eu(url_empleo))
            elif plataforma == "lever":
                ofertas = consultar_lever(slug)
            else:
                ofertas = consultar_ashby(slug)
        except urllib.error.HTTPError as e:
            errores.append((empresa, f"HTTP {e.code} (¿cambiaron de ATS o el slug '{slug}' ya no es válido?)"))
            continue
        except Exception as e:
            errores.append((empresa, str(e)))
            continue

        for oferta in ofertas:
            if not coincide_titulo(oferta["titulo"]):
                continue
            if not oferta["url"] or oferta["url"] in vistas:
                continue
            nuevas.append(
                {
                    "empresa": empresa,
                    "titulo": oferta["titulo"],
                    "url": oferta["url"],
                    "plataforma": plataforma,
                }
            )
            registrar_vista(vistas_writer, empresa, oferta["titulo"], oferta["url"])
            vistas.add(oferta["url"])

        time.sleep(0.4)  # no machacar las APIs

    vistas_f.close()

    with open(SALIDA, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["empresa", "titulo", "url", "plataforma"])
        w.writeheader()
        for n in nuevas:
            w.writerow(n)

    print(f"Empresas consultadas por API: {revisadas}")
    print(f"Ofertas nuevas de diseño encontradas: {len(nuevas)} -> {SALIDA.name}")

    if sin_slug:
        unicas = sorted(set(sin_slug))
        print(
            f"\n{len(unicas)} empresas marcadas greenhouse/lever/ashby pero con un "
            f"dominio propio (ej. empresa.com/careers) donde no se puede extraer "
            f"el slug automáticamente. Para monitorizarlas hay que encontrar su "
            f"enlace directo (boards.greenhouse.io/..., jobs.lever.co/... o "
            f"jobs.ashbyhq.com/...) y actualizar web_empleo con ese enlace:"
        )
        for e in unicas:
            print(f"  - {e}")

    if errores:
        print(f"\n{len(errores)} empresas con error al consultar (revisar a mano):")
        for empresa, err in errores:
            print(f"  - {empresa}: {err}")


if __name__ == "__main__":
    main()
