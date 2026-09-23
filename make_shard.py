#!/usr/bin/env python3
from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
from openpyxl import load_workbook, Workbook


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default="watchlist_consolidado.xlsx")
    ap.add_argument("--output", default="shard_watchlist.xlsx")
    ap.add_argument("--index", type=int, required=True, help="0-based shard index")
    ap.add_argument("--count", type=int, required=True, help="total number of shards")
    args = ap.parse_args()

    if args.index < 0 or args.index >= args.count:
        raise SystemExit("invalid shard index/count")

    src = load_workbook(args.input, read_only=True, data_only=True)
    ws = src["watchlist"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("empty watchlist")

    header, data = rows[0], rows[1:]
    selected = [row for i, row in enumerate(data) if i % args.count == args.index]

    out = Workbook()
    ows = out.active
    ows.title = "watchlist"
    ows.append(list(header))
    for row in selected:
        ows.append(list(row))
    out.save(args.output)

    print(f"Shard {args.index + 1}/{args.count}: {len(selected)} companies -> {args.output}")


if __name__ == "__main__":
    main()
